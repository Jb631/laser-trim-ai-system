"""A file that failed to read before is not offered again (2026-09-17).

James: "i dont want to keep processing repeat files. i just want to process
new stuff."

Two populations came back on EVERY run of the work share:

  A. 67 output-smoothness exports (`*_OS_*`). `detect_file_type` routes them to
     the smoothness parser on the name, the parser finds no usable columns and
     returns no tracks, and `_process_smoothness_file` built an error result
     and returned — recording NOTHING anywhere. The scan offered them as new
     forever. On 2026-09-17: 214 files offered, 179 opened, 110 verdicts; 67
     were these.
  B. 111 old trim exports (8856, 8888, 6952, 8204-3, 8232-1, 8094-2, 8914) the
     parser cannot read. These DO get a `processed_files` row, but with
     success=False, and `_load_processed_hashes` loads only success=True — so
     every run logged "N new, 111 retrying earlier errors".

The rule now: a file that failed, and has not changed on disk since, is not
offered again. The escape hatch is explicit and named on screen — Settings →
Retry unreadable files — for after a parser upgrade.

The mechanism is deliberately the EXISTING per-path skip marker (2026-09-14):
success=True, analysis_id=NULL, a synthetic `skip:` file_hash, and an
`error_message`. What is new is that a failure marker's reason starts with
`UNREADABLE_PREFIX`, and that prefix is the discriminator the retry button
scopes on. "error_message IS NOT NULL" is NOT that discriminator, however much
it looks like one: `mark_file_skipped` appends "(content sha256=…)" to every
marker it writes, and the duplicate markers say "same content as
final_test_results id N" — so a NOT-NULL reset would re-offer the non-trim
junk and the duplicates too. (On the work database today all 8,114 existing
markers carry NULL, which is exactly why the trap is invisible there.)

Transient failures — a locked workbook, a dropped share, a half-written export
— are NEVER marked. They are the one class where "do not offer this again"
would lose real data.
"""

import hashlib
import os
from datetime import datetime
from pathlib import Path

import pytest

from laser_trim_analyzer.core.models import AnalysisStatus
from laser_trim_analyzer.database.manager import DatabaseManager
from laser_trim_analyzer.database.models import (
    UNREADABLE_PREFIX, ProcessedFile)

SAMPLES = Path(__file__).resolve().parents[1] / "Work Files" / "Sample_Base_2026-04-10"


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------

def _db(tmp_path, name="failed_markers.db") -> DatabaseManager:
    return DatabaseManager(tmp_path / name)


class _Row:
    """A detached snapshot — ORM instances expire when the session closes."""

    def __init__(self, r):
        self.file_path = r.file_path
        self.filename = r.filename
        self.file_hash = r.file_hash
        self.file_size = r.file_size
        self.file_modified_date = r.file_modified_date
        self.error_message = r.error_message
        self.success = r.success
        self.analysis_id = r.analysis_id


def _rows(db):
    """Every processed_files row, grouped by path (a path may hold two)."""
    out = {}
    with db.session() as s:
        for r in s.query(ProcessedFile).all():
            out.setdefault(r.file_path, []).append(_Row(r))
    return out


def _markers(db):
    """Only the skip-marker rows, keyed by path."""
    return {p: [r for r in rs if r.analysis_id is None and r.success]
            for p, rs in _rows(db).items()}


def _failure_marker(db, path):
    for r in _markers(db).get(str(path), []):
        if (r.error_message or "").startswith(UNREADABLE_PREFIX):
            return r
    return None


def _processor_loaded_from(db, monkeypatch):
    """A fresh Processor whose caches are loaded from `db` — a NEXT RUN."""
    from laser_trim_analyzer.config import Config
    from laser_trim_analyzer.core.processor import Processor
    import laser_trim_analyzer.database as db_mod

    monkeypatch.setattr(db_mod, "get_database", lambda *a, **k: db)
    proc = Processor(Config(), use_ml=False)
    proc._load_processed_hashes()
    return proc


def _classify_next_run(db, monkeypatch, path):
    """What a fresh run's in-memory scan says about `path`."""
    proc = _processor_loaded_from(db, monkeypatch)
    st = Path(path).stat()
    proc._disk_stats = {str(path): (st.st_size, st.st_mtime)}
    return proc._classify_scan(Path(path))


def _unreadable_os_workbook(path: Path) -> None:
    """An output-smoothness export the parser cannot make tracks from.

    The work-share signature, reproduced exactly: the generic parser logs
    "Generic parser found no usable columns in sheet 'Sheet1'
    (pos_col='Electrical Angle* + 280 min:', smooth_cols=[])" and returns no
    tracks, so `_process_smoothness_file` raises "Smoothness parser returned
    no tracks for …".
    """
    import openpyxl
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.cell(row=1, column=1, value="Electrical Angle* + 280 min:")
    for i in range(2, 8):
        ws.cell(row=i, column=1, value=float(i))
    wb.save(path)


def _os_path(tmp_path) -> Path:
    return (tmp_path / "Test Station" / "6581"
            / "6581-sn330_-65_OS_load_1-2-2017_10-00-00 AM.xlsx")


def _a_trim_sample() -> Path:
    for p in sorted((SAMPLES / "LTS").rglob("*.xls*")):
        if p.is_file() and not p.name.startswith("~$"):
            return p
    pytest.skip("no local trim sample available")


def _error_result(proc, path):
    """The ERROR AnalysisResult a trim file that cannot be parsed produces."""
    return proc._create_error_result(
        proc._create_minimal_metadata(path), "Could not find data start", 0.0)


# ---------------------------------------------------------------------------
# (a) population A — a smoothness file that raises is recorded with its reason
# ---------------------------------------------------------------------------

def test_smoothness_failure_is_recorded_with_its_reason(tmp_path, monkeypatch):
    p = _os_path(tmp_path)
    _unreadable_os_workbook(p)
    db = _db(tmp_path)
    proc = _processor_loaded_from(db, monkeypatch)

    result = proc.process_file(p, generate_plots=False)
    # The result itself is unchanged: still an ERROR, still flagged smoothness
    # so nothing tries to save it as a trim record.
    assert result.overall_status is AnalysisStatus.ERROR
    assert getattr(result, "file_type", None) == "smoothness"

    row = _failure_marker(db, p)
    assert row is not None, (
        "nothing was recorded — this is the bug: 67 of these were offered, "
        "opened and thrown away on every single run")
    assert "returned no tracks" in row.error_message
    assert row.file_hash.startswith("skip:")
    assert row.success is True and row.analysis_id is None
    assert row.file_size == p.stat().st_size


def test_smoothness_failure_is_not_offered_next_run(tmp_path, monkeypatch):
    p = _os_path(tmp_path)
    _unreadable_os_workbook(p)
    db = _db(tmp_path)
    _processor_loaded_from(db, monkeypatch).process_file(p, generate_plots=False)

    # No change to `_load_processed_hashes` or to the scan was needed: the
    # marker is success=True, so the existing success-only load picks it up and
    # the existing stat fast-path skips it from memory alone.
    assert _classify_next_run(db, monkeypatch, p) == "processed"


# ---------------------------------------------------------------------------
# (b) the same for a final-test parse failure
# ---------------------------------------------------------------------------

def _ft_path(tmp_path) -> Path:
    p = (tmp_path / "Test Station" / "1844205"
         / "1844205-sn80A_VO_8-19-2026_7-36-16 AM.xlsx")
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_bytes(b"not a workbook, but present and readable")
    return p


def test_final_test_failure_is_recorded_and_skipped(tmp_path, monkeypatch):
    p = _ft_path(tmp_path)
    db = _db(tmp_path)
    proc = _processor_loaded_from(db, monkeypatch)

    def boom(*a, **k):
        raise ValueError("Could not find data start")

    monkeypatch.setattr(proc.final_test_parser, "parse_file", boom)
    result = proc._process_final_test_file(p, 0.0)
    assert result.overall_status is AnalysisStatus.ERROR
    assert getattr(result, "file_type", None) == "final_test"

    row = _failure_marker(db, p)
    assert row is not None, "a final-test file that cannot be parsed vanished"
    assert "Could not find data start" in row.error_message
    assert "ValueError" in row.error_message
    assert _classify_next_run(db, monkeypatch, p) == "processed"


# ---------------------------------------------------------------------------
# (c) population B — a trim ERROR keeps its success=False row AND gets a marker
# ---------------------------------------------------------------------------

def test_trim_error_keeps_its_row_and_gains_a_marker(tmp_path, monkeypatch):
    import shutil
    src = _a_trim_sample()
    p = tmp_path / "LTS" / src.name
    p.parent.mkdir(parents=True)
    shutil.copy2(src, p)

    db = _db(tmp_path)
    proc = _processor_loaded_from(db, monkeypatch)
    db.save_analysis(_error_result(proc, p))

    rows = _rows(db)[str(p)]
    assert len(rows) == 2, (
        "one path, two rows: the ERROR row keyed by content hash and the "
        f"marker keyed by path — got {[(r.file_hash[:12], r.success) for r in rows]}")
    error_row = [r for r in rows if r.success is False]
    marker = [r for r in rows if r.success is True]
    assert len(error_row) == 1 and len(marker) == 1
    # Unchanged: ERROR results are still recorded success=False with their
    # analysis, which is what every existing query and the cleanup tools read.
    assert error_row[0].analysis_id is not None
    assert not error_row[0].file_hash.startswith("skip:")
    assert marker[0].analysis_id is None
    assert marker[0].file_hash.startswith("skip:")
    assert marker[0].error_message.startswith(UNREADABLE_PREFIX)
    assert "Could not find data start" in marker[0].error_message


def test_trim_error_is_not_offered_next_run(tmp_path, monkeypatch):
    import shutil
    src = _a_trim_sample()
    p = tmp_path / "LTS" / src.name
    p.parent.mkdir(parents=True)
    shutil.copy2(src, p)

    db = _db(tmp_path)
    proc = _processor_loaded_from(db, monkeypatch)
    db.save_analysis(_error_result(proc, p))

    assert _classify_next_run(db, monkeypatch, p) == "processed", (
        "this is the '111 retrying earlier errors' line, every run since 2013")


def test_trim_error_row_is_still_a_retryable_error_row(tmp_path, monkeypatch):
    """The marker must not have been implemented by flipping `success`.

    `is_success` and `_error_basenames` decide how the app TALKS about a file;
    the marker decides whether it is OFFERED. Conflating them would hide the
    error from the cleanup tools and from the scan message.
    """
    import shutil
    src = _a_trim_sample()
    p = tmp_path / "LTS" / src.name
    p.parent.mkdir(parents=True)
    shutil.copy2(src, p)

    db = _db(tmp_path)
    proc = _processor_loaded_from(db, monkeypatch)
    db.save_analysis(_error_result(proc, p))

    proc2 = _processor_loaded_from(db, monkeypatch)
    assert p.name in proc2._error_basenames


# ---------------------------------------------------------------------------
# (d) transient failures are never marked
# ---------------------------------------------------------------------------

def test_transient_classifier():
    from laser_trim_analyzer.core.processor import Processor

    transient = Processor._is_transient_failure
    assert transient(PermissionError(13, "Permission denied"))
    assert transient(FileNotFoundError(2, "No such file or directory"))
    assert transient(TimeoutError("timed out"))
    assert transient(OSError("[Errno 51] Network is unreachable"))
    # …and the message-only form, which is all the manager has to go on.
    assert transient("[Errno 13] Permission denied: 'X.xls'")
    assert transient("[WinError 32] The process cannot access the file because "
                     "it is being used by another process")

    # A parse failure is NOT transient, however it is spelled — including the
    # OSError-shaped ones, which is why the permanent taxonomy wins.
    assert not transient(ValueError("Could not find data start"))
    assert not transient(ValueError("positions not monotonically increasing"))
    assert not transient(Exception("limit columns are not a +/- band"))
    assert not transient(OSError("Excel file format cannot be determined"))
    assert not transient("Smoothness parser returned no tracks for X.xlsx")


def test_a_locked_smoothness_file_is_not_marked(tmp_path, monkeypatch):
    p = _os_path(tmp_path)
    _unreadable_os_workbook(p)
    db = _db(tmp_path)
    proc = _processor_loaded_from(db, monkeypatch)

    def locked(*a, **k):
        raise PermissionError(13, "Permission denied", str(p))

    monkeypatch.setattr(proc.smoothness_parser, "parse_file", locked)
    proc._process_smoothness_file(p, 0.0)

    assert _failure_marker(db, p) is None, (
        "a workbook someone had open in Excel would never be read again")
    assert _classify_next_run(db, monkeypatch, p) == "new"


def test_a_locked_trim_file_is_not_marked(tmp_path, monkeypatch):
    """The manager side of the same rule: it only has the reason TEXT."""
    import shutil
    src = _a_trim_sample()
    p = tmp_path / "LTS" / src.name
    p.parent.mkdir(parents=True)
    shutil.copy2(src, p)

    db = _db(tmp_path)
    proc = _processor_loaded_from(db, monkeypatch)
    result = proc._create_error_result(
        proc._create_minimal_metadata(p),
        f"[Errno 13] Permission denied: '{p}'", 0.0)
    db.save_analysis(result)

    assert _failure_marker(db, p) is None
    assert _classify_next_run(db, monkeypatch, p) == "new"


# ---------------------------------------------------------------------------
# (e) a marked file that CHANGED on disk is offered again
# ---------------------------------------------------------------------------

def test_a_changed_file_is_offered_again(tmp_path, monkeypatch):
    p = _os_path(tmp_path)
    _unreadable_os_workbook(p)
    db = _db(tmp_path)
    _processor_loaded_from(db, monkeypatch).process_file(p, generate_plots=False)
    assert _classify_next_run(db, monkeypatch, p) == "processed"

    # Re-exported on the share (or: it was half-written the first time).
    p.write_bytes(p.read_bytes() + b"\0" * 4096)
    st = p.stat()
    os.utime(p, (st.st_atime, st.st_mtime + 3600))

    proc = _processor_loaded_from(db, monkeypatch)
    st = p.stat()
    proc._disk_stats = {str(p): (st.st_size, st.st_mtime)}
    assert proc._classify_scan(p) == "needs_hash"
    assert proc._is_processed(p) is False, (
        "a file whose bytes changed must be read again — this is what makes a "
        "half-written export self-heal without anyone pressing anything")


# ---------------------------------------------------------------------------
# (f) the retry clears failure markers ONLY
# ---------------------------------------------------------------------------

def _plain_skip(db, path: Path, reason=None):
    """What the processor records for a non-trim file (no reason of its own)."""
    st = path.stat()
    db.mark_file_skipped(
        filename=path.name, file_path=str(path),
        file_hash=hashlib.sha256(path.read_bytes()).hexdigest(),
        file_size=st.st_size,
        file_modified_date=datetime.fromtimestamp(st.st_mtime),
        error_message=reason)


def test_retry_clears_only_the_files_that_failed_to_read(tmp_path, monkeypatch):
    junk = tmp_path / "parameters.xlsx"
    junk.write_bytes(b"a parameter workbook, not test data")
    dupe = tmp_path / "Copy of 8084-sn1.xls"
    dupe.write_bytes(b"a duplicate export")
    os_file = _os_path(tmp_path)
    _unreadable_os_workbook(os_file)

    db = _db(tmp_path)
    _plain_skip(db, junk)                                   # non-trim: no reason
    _plain_skip(db, dupe, reason="same content as final_test_results id 7")
    _processor_loaded_from(db, monkeypatch).process_file(os_file,
                                                        generate_plots=False)

    assert db.count_skipped_files() == 3
    # The trap, made explicit: BOTH of the innocent markers carry a non-NULL
    # error_message, because `mark_file_skipped` records the content hash in
    # it. "error_message IS NOT NULL" would clear all three.
    marks = _markers(db)
    assert marks[str(junk)][0].error_message
    assert marks[str(dupe)][0].error_message
    assert db.count_failed_file_markers() == 1, (
        "the count must see the read failure and only the read failure")

    cleared = db.reset_failed_file_markers()
    assert cleared == 1
    assert db.count_failed_file_markers() == 0
    left = set(_markers(db))
    assert str(junk) in left and str(dupe) in left, (
        "retrying read failures must not re-offer the 8,114 non-trim files or "
        "the duplicates — they never failed to read")
    assert str(os_file) not in left or not _markers(db)[str(os_file)]

    # And the point of the button: the file is offered again.
    assert _classify_next_run(db, monkeypatch, os_file) == "new"


def test_reset_skipped_files_still_clears_everything(tmp_path, monkeypatch):
    """The blunt instrument next to it is unchanged."""
    junk = tmp_path / "parameters.xlsx"
    junk.write_bytes(b"a parameter workbook")
    os_file = _os_path(tmp_path)
    _unreadable_os_workbook(os_file)
    db = _db(tmp_path)
    _plain_skip(db, junk)
    _processor_loaded_from(db, monkeypatch).process_file(os_file,
                                                        generate_plots=False)
    assert db.reset_skipped_files() == 2
    assert db.count_skipped_files() == 0


def test_a_file_that_later_succeeds_stops_being_counted(tmp_path, monkeypatch):
    """A marker is a statement about the bytes that failed, not about the path.

    Without this, a file that failed once and was then re-exported correctly
    would keep HOME saying "1 file is being skipped" forever.
    """
    import shutil
    src = _a_trim_sample()
    p = tmp_path / "LTS" / src.name
    p.parent.mkdir(parents=True)
    shutil.copy2(src, p)

    db = _db(tmp_path)
    proc = _processor_loaded_from(db, monkeypatch)
    db.save_analysis(_error_result(proc, p))
    assert db.count_failed_file_markers() == 1

    good = proc.process_file(p, generate_plots=False)
    assert good.overall_status is not AnalysisStatus.ERROR
    db.save_analysis(good)
    assert db.count_failed_file_markers() == 0, (
        "the file reads fine now; nothing is being skipped")


def test_deleting_the_error_record_re_offers_the_file(tmp_path, monkeypatch):
    """`delete_analysis` deletes the processed row "to allow re-processing".

    The marker is keyed by PATH and carries no analysis_id, so it would
    outlive the record and make that promise false.
    """
    import shutil
    src = _a_trim_sample()
    p = tmp_path / "LTS" / src.name
    p.parent.mkdir(parents=True)
    shutil.copy2(src, p)

    db = _db(tmp_path)
    proc = _processor_loaded_from(db, monkeypatch)
    analysis_id = db.save_analysis(_error_result(proc, p))
    assert _classify_next_run(db, monkeypatch, p) == "processed"

    assert db.delete_analysis(analysis_id) is True
    assert _rows(db).get(str(p)) is None, "a row survived the delete"
    assert _classify_next_run(db, monkeypatch, p) == "new"


# ---------------------------------------------------------------------------
# (g) two rows for one path break nothing
# ---------------------------------------------------------------------------

def test_two_rows_for_one_path_do_not_break_the_recording_path(tmp_path,
                                                               monkeypatch):
    """`_record_processed_file` does `scalar_one_or_none()` BY CONTENT HASH.

    A second row for the same path is only safe because its hash is synthetic:
    `idx_processed_path` is non-unique, UNIQUE(file_hash) is satisfied because
    the two hashes differ, and every content-keyed query still matches exactly
    one row. If the marker ever carried the real hash this raises
    MultipleResultsFound and the whole trim save path dies.
    """
    import shutil
    from sqlalchemy import select
    from laser_trim_analyzer.utils.hashing import calculate_file_hash

    src = _a_trim_sample()
    p = tmp_path / "LTS" / src.name
    p.parent.mkdir(parents=True)
    shutil.copy2(src, p)

    db = _db(tmp_path)
    proc = _processor_loaded_from(db, monkeypatch)
    db.save_analysis(_error_result(proc, p))
    assert len(_rows(db)[str(p)]) == 2

    content_hash = calculate_file_hash(p)
    with db.session() as s:
        got = s.execute(
            select(ProcessedFile.id)
            .where(ProcessedFile.file_hash == content_hash)
        ).scalar_one_or_none()          # must not raise MultipleResultsFound
    assert got is not None

    # is_file_processed is content-keyed too: the ERROR row is success=False
    # and the marker is invisible to it, so the answer is still "no".
    assert db.is_file_processed(p) is False

    # And saving the same file again (the re-offer after a Retry) still works.
    db.save_analysis(_error_result(proc, p))
    assert len(_rows(db)[str(p)]) == 2, "a second failure must not add a row"

    # The stat heal, also content-keyed, still cannot see the marker.
    healed = db.update_processed_file_stats(
        [(content_hash, 4242, datetime(2026, 9, 17))])
    assert healed["processed_files"] == 1
    marker = _failure_marker(db, p)
    assert marker.file_size == p.stat().st_size, "the heal rewrote a marker"


# ---------------------------------------------------------------------------
# never silent: the line HOME shows and the button Settings offers
# ---------------------------------------------------------------------------

def test_the_notice_wording():
    from laser_trim_analyzer.core.ingest_run import unreadable_notice

    assert unreadable_notice(0) == ""
    assert unreadable_notice(-1) == ""
    line = unreadable_notice(178)
    assert "178 files" in line
    assert "failed to read" in line
    assert "Settings → Retry unreadable files" in line
    assert "1 file is" in unreadable_notice(1)


def test_unreadable_count_survives_an_old_database():
    from laser_trim_analyzer.core.ingest_run import unreadable_count

    class _Broken:
        def count_failed_file_markers(self):
            raise RuntimeError("no such column")

    assert unreadable_count(_Broken()) == 0


def test_home_shows_the_line_only_when_there_is_something_to_say(make_app,
                                                                 monkeypatch):
    app = make_app()
    page = app.page_container.get_page("home")

    monkeypatch.setattr(type(app.db), "count_failed_file_markers",
                        lambda self: 0, raising=False)
    page.reload_now()
    app.update_idletasks()
    assert page._unreadable_label.cget("text") == ""
    assert not page._unreadable_label.winfo_ismapped()

    monkeypatch.setattr(type(app.db), "count_failed_file_markers",
                        lambda self: 178, raising=False)
    page.reload_now()
    app.update_idletasks()
    said = page._unreadable_label.cget("text")
    assert "178 files" in said and "Retry unreadable files" in said


def test_settings_offers_the_retry_and_reports_what_it_cleared(make_app,
                                                               monkeypatch):
    import customtkinter as ctk
    from tkinter import messagebox
    from laser_trim_analyzer.gui.v6.sections import database_cleanup

    app = make_app()
    monkeypatch.setattr(type(app.db), "count_failed_file_markers",
                        lambda self: 178, raising=False)
    cleared = []
    monkeypatch.setattr(type(app.db), "reset_failed_file_markers",
                        lambda self: cleared.append(1) or 178, raising=False)
    asked = []
    monkeypatch.setattr(messagebox, "askyesno",
                        lambda title, text: asked.append(text) or True)

    frame = ctk.CTkFrame(app)
    try:
        database_cleanup.build_database_cleanup_section(frame, app.theme, app)
        button = _find_button(frame, "Retry unreadable files")
        assert button is not None, "Settings offers no way back"
        button.invoke()
        _settle(app, until=lambda: cleared and _said(frame, "178"))
        assert asked and "178" in asked[0], asked
        assert cleared == [1]
        assert _said(frame, "178"), [w.cget("text") for w in _walk(frame)
                                     if isinstance(w, ctk.CTkLabel)]
    finally:
        frame.destroy()


def _walk(widget):
    yield widget
    for child in widget.winfo_children():
        yield from _walk(child)


def _find_button(root, text):
    import customtkinter as ctk
    for w in _walk(root):
        if isinstance(w, ctk.CTkButton) and w.cget("text") == text:
            return w
    return None


def _said(root, text):
    import customtkinter as ctk
    return any(isinstance(w, ctk.CTkLabel) and text in (w.cget("text") or "")
               for w in _walk(root))


def _settle(app, until, seconds=5.0):
    """Let the worker threads finish and their posted callbacks run."""
    import time
    deadline = time.monotonic() + seconds
    while time.monotonic() < deadline:
        app.update()
        if until():
            return
        time.sleep(0.02)
