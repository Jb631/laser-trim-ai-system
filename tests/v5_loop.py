"""V5's loop over every kind of file the ingest meets, and the rows it leaves (ingest-speed Task 9).

Task 9 takes the worker's writes out of the worker: the pool now runs `analyse_path`, which
touches no database and returns an Outcome, and the consumer applies it -- through a writer when
one is given (V6), at once through the public methods when none is (V5). The promise is the one
every save-path task here keeps: not one stored value changes. So this is V5's own loop --
`process_batch` with no writer, `db.save_analysis(result)` for every result it yields, exactly
as gui/pages/process.py does it -- over one scenario, and `fixtures/v5_loop_golden.json` holds
the rows it left at f232528, BEFORE the worker's writes moved. Regenerate only on purpose:

    LTA_REGENERATE_SAVE_ROWS=1 python -m pytest tests/test_worker_outcomes.py -k golden

THE SCENARIO (every value not read from a fixture is invented):
  * trims -- System A, System B, a TrimVolts file (laser 1's increment curves), a trim-setup
    file, a two-track file, and a no-cut template (an UNTRIMMED result);
  * final tests -- every format the repository holds: Format 1 x4, Format 3's three tracks,
    Format 4 without a serial (its save is refused: a PERMANENT failure, so a marker with no
    reason), and a synthetic Format 2 `Rout_` workbook;
  * smoothness -- a synthetic Betatronix export (saved), and one the parser finds no tracks in
    (a marker with its reason);
  * not test data -- a file named as noise (non_trim: a marker, nothing yielded), and a
    parameter workbook named like a trim (NonTrimWorkbookError: the same);
  * final-test failures -- junk bytes (a PERMANENT failure: a marker with no reason), a
    workbook with no data sheet (a header-only row), and a truncated workbook (an exception:
    a marker WITH its reason);
  * trim failures -- junk bytes the detector cannot open (non_trim: a marker), junk bytes the
    trim parser meets (a PERMANENT failure: a marker AND the ERROR result the loop saves), a
    truncated workbook (an exception: the ERROR result alone), and sheets with no sweep in them
    ("No valid track data found").
"""
from __future__ import annotations

import os
import shutil
from pathlib import Path
from typing import List, Tuple

import save_rows
from save_rows import FIXTURES, MTIME, _pinned_copy

TABLES = ("analysis_results", "track_results", "trim_passes", "trim_setup", "processed_files",
          "final_test_results", "final_test_tracks", "smoothness_results", "smoothness_tracks")
GOLDEN = FIXTURES / "v5_loop_golden.json"


def _pinned(path: Path) -> Path:
    os.utime(path, (MTIME, MTIME))
    return path


def _betatronix_os(path: Path, *, max_dev: float, spec: float, result: str) -> Path:
    """A synthetic Betatronix Output Smoothness export (invented values)."""
    import openpyxl
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Data"
    header = ["Model Parameters", "Model Number", "9991", None, "Time (s)",
              "Filtered Volts (V)", None, "Max. Deviation (V)", "Max. Spec. Dev. (V)", "Result"]
    for col, value in enumerate(header, start=1):
        ws.cell(row=1, column=col, value=value)
    for i in range(12):
        ws.cell(row=2 + i, column=5, value=round(0.1 * i, 3))
        ws.cell(row=2 + i, column=6, value=round(0.002 * ((i * 7) % 5) - 0.004, 5))
    ws.cell(row=2, column=8, value=max_dev)
    ws.cell(row=2, column=9, value=spec)
    ws.cell(row=2, column=10, value=result)
    wb.save(path)
    return _pinned(path)


def _parameter_workbook(path: Path) -> Path:
    """A laser's parameter/report workbook named like test data: sheets, no sweep."""
    import openpyxl
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    wb.active.title = "Model Parameters"
    wb.active.cell(row=1, column=1, value="Model")
    wb.create_sheet("Trim Parameters").cell(row=1, column=1, value="Laser Cut Length")
    wb.save(path)
    return _pinned(path)


def _sheets_only(path: Path, sheets) -> Path:
    """A workbook with these sheet names and nothing a parser can use (invented)."""
    import openpyxl
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    wb.active.title = sheets[0]
    for name in sheets[1:]:
        wb.create_sheet(name)
    wb.active.cell(row=1, column=1, value="invented")
    wb.save(path)
    return _pinned(path)


def build_v5_scenario(root: Path) -> List[Path]:
    """Every file, in the order V5 would be handed them (the walk's order is the caller's)."""
    from side_writes import FT_FIXTURES, ROUT
    from test_failed_file_markers import _unreadable_os_workbook
    from test_ft_polyfit_degenerate import clean_series, write_format2
    from test_no_cut_template import _make_workbook

    inp = root / "in"
    trims = [_pinned_copy(FIXTURES / "trim" / n, inp / "laser" / n)
             for n in ("dlts_8232-1_242.xls", "lts_8232-1_193.xls", "dlts_8074_18.xls",
                       "dlts_7553_10B.xls")]
    trims.append(_pinned_copy(FIXTURES / "trimvolts" / "lts_8340-1_32.xls",
                              inp / "laser" / "lts_8340-1_32.xls"))
    trims.append(_pinned_copy(FIXTURES / "trim_setup" / "8251-1_29_template_updated.xls",
                              inp / "laser" / "8251-1_29_template_updated.xls"))
    nocut = _pinned(_make_workbook(inp / "laser" / "9999_12_TA_Test Data_1-2-2026_9-00 AM.xlsx",
                                   lin_error="template"))
    # Under a "Test Station" folder, as on the share: the Format 4 file carries no final-test
    # filename pattern, and the detector routes it by that folder.
    station = inp / "Test Station" / "ft"
    fts = [_pinned_copy(FIXTURES / "final_test" / n, station / n) for n in FT_FIXTURES]
    rout = _pinned(write_format2(station / ROUT, *clean_series()))
    ft_junk = station / "9996-sn1_1-2-2026_9-00 AM.xls"          # a final test nothing can read
    ft_junk.write_bytes(b"invented: not a workbook at all")
    _pinned(ft_junk)
    os_ok = _betatronix_os(inp / "os" / "9991-sn4_OS_2-3-2026_10-00-00 AM.xlsx",
                           max_dev=0.0031, spec=0.0050, result="PASSED")
    os_bad = inp / "os" / "9991-sn5_OS_2-3-2026_10-05-00 AM.xlsx"
    _unreadable_os_workbook(os_bad)
    _pinned(os_bad)
    noise = inp / "junk" / "9993_noise_capture.xls"               # not test data, by name
    noise.parent.mkdir(parents=True, exist_ok=True)
    noise.write_bytes(b"invented: an oscilloscope capture")
    _pinned(noise)
    params = _parameter_workbook(inp / "laser" / "9994_3_TA_Test Data_1-2-2026_9-00 AM.xlsx")
    unopenable = inp / "laser" / "9992_7_TA_Test Data_1-2-2026_9-00 AM.xls"  # the detector cannot
    unopenable.write_bytes(b"invented: not a workbook at all")               # open it: non_trim
    _pinned(unopenable)
    # "Trimmed" in the name routes to the trim parser without opening the file first.
    permanent = inp / "laser" / "9997_2_TA_Test Data_1-2-2026_9-00 AMTrimmed Correct.xls"
    permanent.write_bytes(b"invented: not a workbook at all")    # "format cannot be determined"
    _pinned(permanent)
    truncated = inp / "laser" / "9997_3_TA_Test Data_1-2-2026_9-00 AMTrimmed Correct.xls"
    truncated.write_bytes((FIXTURES / "trim" / "lts_8232-1_194.xls").read_bytes()[:20000])
    _pinned(truncated)                                            # an exception, NOT permanent
    no_tracks = _sheets_only(inp / "laser" / "9998_5_TA_Test Data_1-2-2026_9-00 AM.xlsx",
                             ["SEC1 TRK1 0", "SEC1 TRK1 TRM"])    # "No valid track data found"
    ft_header_only = _sheets_only(station / "9996-sn2_1-2-2026_9-00 AM.xlsx", ["Nothing"])
    ft_truncated = station / "9996-sn4_1-2-2026_9-00 AM.xls"
    ft_truncated.write_bytes((FIXTURES / "final_test" / FT_FIXTURES[0]).read_bytes()[:20000])
    _pinned(ft_truncated)
    return (trims + [nocut] + fts + [rout, ft_junk, ft_header_only, ft_truncated, os_ok, os_bad,
                                     noise, params, unopenable, permanent, truncated, no_tracks])


def v5_loop(db, paths, root: Path, *, parallel: bool, writer=None) -> List[Tuple[str, str, str]]:
    """V5's loop (gui/pages/process.py): process_batch with no writer, db.save_analysis for every
    result it yields. Returns what V5 saw: (filename, file type, status) per yielded result."""
    from laser_trim_analyzer.config import Config
    from laser_trim_analyzer.core.processor import Processor
    config = Config()
    config.processing.turbo_mode_threshold = 1 if parallel else 10 ** 9
    proc = Processor(config=config, use_ml=False)
    proc.ml_storage_path = root / "no_ml_models"      # never the CWD's data/ml_models
    seen = []
    kwargs = {} if writer is None else {"writer": writer}
    for result in proc.process_batch([Path(p) for p in paths], incremental=False, **kwargs):
        seen.append((result.metadata.filename, getattr(result, "file_type", "trim"),
                     result.overall_status.value))
        db.save_analysis(result)
    return seen


def v5_snapshot(db_path: Path, root: Path, seen) -> dict:
    """The rows, normalised as save_rows does -- plus the one clock V5's rows carry: an ERROR
    result built from no metadata is dated `datetime.now()` (_create_minimal_metadata)."""
    snap = {"seen": [list(s) for s in seen], "tables": save_rows.dump_rows(db_path, root, TABLES)}
    for row in snap["tables"]["analysis_results"]:
        if row["overall_status"] == "ERROR" and row["model"] == "Unknown":
            row["file_date"] = "<now>"
            row["unit_id"] = "<now>"
    return snap


def order_free(tables: dict) -> dict:
    """The same rows with every id replaced by the natural key of the row it points at, in a
    stable order. Two runs that stored the same rows in a different ORDER -- the batch writer
    commits a failed save's marker with the NEXT batch, a pool completes files in any order --
    compare equal; nothing else does."""
    import json

    def natural(row, *cols):
        return json.dumps([row.get(c) for c in cols], default=str)

    ak = {r["id"]: natural(r, "filename", "file_date", "model", "serial")
          for r in tables.get("analysis_results", [])}
    tk = {r["id"]: json.dumps([ak.get(r["analysis_id"]), r["track_id"]])
          for r in tables.get("track_results", [])}
    fk = {r["id"]: natural(r, "filename", "file_date", "model", "serial")
          for r in tables.get("final_test_results", [])}
    sk = {r["id"]: natural(r, "filename", "file_date", "model", "serial")
          for r in tables.get("smoothness_results", [])}
    refs = {"analysis_id": ak, "linked_trim_id": ak, "track_result_id": tk,
            "final_test_id": fk, "smoothness_id": sk}
    out = {}
    for table, rows in tables.items():
        conv = [{c: (refs[c].get(v, f"<dangling {v}>") if c in refs and v is not None else v)
                 for c, v in r.items() if c != "id"} for r in rows]
        out[table] = sorted(conv, key=lambda r: json.dumps(r, sort_keys=True, default=str))
    return out
