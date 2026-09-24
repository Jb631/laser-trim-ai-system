"""Task 3b: reading a stored record back never invents a verdict or a number.

`database/manager.py::_map_db_to_track` turns stored rows back into `TrackData`
for `db.get_analysis()` / `db.get_historical_data()`. Before this fix it turned a
NULL `linearity_pass` into False, recomputed a NULL `sigma_pass` from the
analyzer's 999.999 saturation marker (999.999 <= threshold -> False), and turned
a NULL `final_linearity_error_shifted` into 0.0 -- so a track that FAILED
PROCESSING (core.model_stats.failed_processing) came back looking like a graded
FAIL. CLAUDE.md, "What must never be stored": a blank measurement is ungraded,
never 0.0; a record that failed processing is not a measurement; a failure must
never look like a result.

This file pins the read-back itself, plus every V5 GUI / export consumer that
reads linearity_pass/sigma_pass/linearity_error/sigma_gradient off a TrackData:
- gui/pages/analyze.py: _display_track_chart, _display_metrics,
  _plot_error_vs_position_export, _plot_metrics_summary, _plot_status_display
- gui/pages/export.py: _plot_error_vs_position_export, _plot_metrics_summary,
  _plot_status_display
- export/excel.py: _create_tracks_sheet (export_single_result)

gui/pages/quality_health.py and gui/v6/widgets/unit_chart_modal.py query
DBTrackResult directly (never through _map_db_to_track) and are unaffected --
not covered here. export/excel.py's _create_summary_sheet and
_create_all_results_sheet were already None-safe before this fix and are only
smoke-checked in passing.
"""
from datetime import datetime
from pathlib import Path

import pytest

from laser_trim_analyzer.core.analyzer import Analyzer
from laser_trim_analyzer.core.models import (
    AnalysisResult,
    AnalysisStatus,
    FileMetadata,
    RiskCategory,
    SystemType,
    TrackData,
)

MODEL = "TESTMODEL-3B"


def _error_track() -> TrackData:
    """The real shape `_create_failed_track` writes (< 10 points -> ERROR):
    status ERROR, sigma_gradient/linearity_error 999.999 (the analyzer's
    saturation marker, not a reading), sigma_pass/linearity_pass NULL."""
    t = Analyzer().analyze_track({
        "track_id": "default",
        "positions": [float(i) for i in range(5)],
        "errors": [0.001] * 5,
        "upper_limits": [0.01] * 5,
        "lower_limits": [-0.01] * 5,
        "travel_length": 1.0,
        "linearity_spec": 0.01,
    })
    assert t.status == AnalysisStatus.ERROR
    assert t.linearity_pass is None and t.sigma_pass is None  # sanity: Task 1's shape
    return t


def _graded_track(status, sigma_gradient, sigma_threshold, sigma_pass,
                   linearity_error, linearity_pass, fail_points=0) -> TrackData:
    """A normally-graded track with real measurement arrays behind its verdict."""
    n = 12
    return TrackData(
        track_id="default",
        status=status,
        travel_length=1.0,
        linearity_spec=0.01,
        sigma_gradient=sigma_gradient,
        sigma_threshold=sigma_threshold,
        sigma_pass=sigma_pass,
        optimal_offset=0.0,
        linearity_error=linearity_error,
        linearity_pass=linearity_pass,
        linearity_fail_points=fail_points,
        position_data=[float(i) for i in range(n)],
        error_data=[0.001] * n,
        upper_limits=[0.01] * n,
        lower_limits=[-0.01] * n,
        risk_category=RiskCategory.LOW,
    )


def _wrap(track: TrackData, status: AnalysisStatus, serial: str) -> AnalysisResult:
    return AnalysisResult(
        metadata=FileMetadata(
            filename=f"{serial}.xls",
            file_path=Path(f"/fake/{serial}.xls"),
            file_date=datetime.now(),
            model=MODEL,
            serial=serial,
            system=SystemType.B,
        ),
        overall_status=status,
        processing_time=0.1,
        tracks=[track],
    )


@pytest.fixture
def db(tmp_path, monkeypatch):
    """A throwaway DatabaseManager with BOTH _db_manager globals injected --
    get_database() (used by several call sites) ignores whatever manager a
    test builds and constructs its own at the config default otherwise."""
    from laser_trim_analyzer.database import manager as mgr
    import laser_trim_analyzer.database as dbpkg
    d = mgr.DatabaseManager(tmp_path / "read_back.db")
    monkeypatch.setattr(mgr, "_db_manager", d, raising=False)
    monkeypatch.setattr(dbpkg, "_db_manager", d, raising=False)
    return d


# ---------------------------------------------------------------------------
# _map_db_to_track itself, via get_analysis() and get_historical_data()
# ---------------------------------------------------------------------------

def test_get_analysis_error_track_has_no_invented_verdict_or_number(db):
    error_id = db.save_analysis(_wrap(_error_track(), AnalysisStatus.ERROR, "ERR-1"))

    err = db.get_analysis(error_id).tracks[0]
    assert err.linearity_pass is None
    assert err.sigma_pass is None
    assert err.linearity_error is None
    assert err.sigma_gradient is None
    assert err.sigma_threshold is None


def test_get_analysis_graded_tracks_are_unchanged(db):
    fail_id = db.save_analysis(_wrap(
        _graded_track(AnalysisStatus.FAIL, 0.05, 0.03, False, 0.02, False, fail_points=3),
        AnalysisStatus.FAIL, "FAIL-1"))
    pass_id = db.save_analysis(_wrap(
        _graded_track(AnalysisStatus.PASS, 0.01, 0.03, True, 0.005, True),
        AnalysisStatus.PASS, "PASS-1"))

    fail = db.get_analysis(fail_id).tracks[0]
    assert fail.linearity_pass is False
    assert fail.sigma_pass is False
    assert fail.linearity_error == pytest.approx(0.02)
    assert fail.sigma_gradient == pytest.approx(0.05)
    assert fail.sigma_threshold == pytest.approx(0.03)

    passed = db.get_analysis(pass_id).tracks[0]
    assert passed.linearity_pass is True
    assert passed.sigma_pass is True
    assert passed.linearity_error == pytest.approx(0.005)
    assert passed.sigma_gradient == pytest.approx(0.01)
    assert passed.sigma_threshold == pytest.approx(0.03)


def test_get_historical_data_error_track_has_no_invented_verdict_or_number(db):
    db.save_analysis(_wrap(_error_track(), AnalysisStatus.ERROR, "ERR-2"))
    db.save_analysis(_wrap(
        _graded_track(AnalysisStatus.FAIL, 0.05, 0.03, False, 0.02, False, fail_points=3),
        AnalysisStatus.FAIL, "FAIL-2"))
    db.save_analysis(_wrap(
        _graded_track(AnalysisStatus.PASS, 0.01, 0.03, True, 0.005, True),
        AnalysisStatus.PASS, "PASS-2"))

    by_serial = {r.metadata.serial: r for r in
                 db.get_historical_data(model=MODEL, days_back=3650, limit=10)}
    assert set(by_serial) == {"ERR-2", "FAIL-2", "PASS-2"}

    err = by_serial["ERR-2"].tracks[0]
    assert err.linearity_pass is None
    assert err.sigma_pass is None
    assert err.linearity_error is None
    assert err.sigma_gradient is None
    assert err.sigma_threshold is None

    fail = by_serial["FAIL-2"].tracks[0]
    assert fail.linearity_pass is False and fail.sigma_pass is False
    assert fail.sigma_gradient == pytest.approx(0.05)

    passed = by_serial["PASS-2"].tracks[0]
    assert passed.linearity_pass is True and passed.sigma_pass is True
    assert passed.sigma_gradient == pytest.approx(0.01)


# ---------------------------------------------------------------------------
# gui/pages/analyze.py (V5 Analyze page)
# ---------------------------------------------------------------------------

class _MetricsStub:
    """Minimal stand-in for AnalyzePage: _display_metrics only touches
    self._update_metrics, never any Tk widget -- no tk_root/make_app needed."""
    def __init__(self):
        self.text = None

    def _update_metrics(self, text):
        self.text = text


def test_analyze_page_display_metrics_error_track_is_not_graded(db):
    from laser_trim_analyzer.gui.pages.analyze import AnalyzePage

    error_id = db.save_analysis(_wrap(_error_track(), AnalysisStatus.ERROR, "ERR-3"))
    analysis = db.get_analysis(error_id)

    stub = _MetricsStub()
    AnalyzePage._display_metrics(stub, analysis)  # must not raise

    assert stub.text is not None
    assert "999.999" not in stub.text
    assert "not graded" in stub.text.lower()
    assert "SIGMA ANALYSIS" not in stub.text  # the numeric block is skipped entirely
    assert "✗ FAIL" not in stub.text     # '✗ FAIL'
    assert "✓ PASS" not in stub.text     # '✓ PASS'


def test_analyze_page_display_metrics_graded_fail_track_still_shows_numbers(db):
    """Regression guard: the failed-processing branch must not swallow a real grade."""
    from laser_trim_analyzer.gui.pages.analyze import AnalyzePage

    fail_id = db.save_analysis(_wrap(
        _graded_track(AnalysisStatus.FAIL, 0.05, 0.03, False, 0.02, False, fail_points=3),
        AnalysisStatus.FAIL, "FAIL-3"))
    analysis = db.get_analysis(fail_id)

    stub = _MetricsStub()
    AnalyzePage._display_metrics(stub, analysis)

    assert "not graded" not in stub.text.lower()
    assert "SIGMA ANALYSIS" in stub.text
    assert "0.050000" in stub.text  # sigma_gradient formatted with :.6f


def test_analyze_page_display_track_chart_error_track_does_not_crash():
    """_display_track_chart touches self.chart/self.current_result -- stub both.
    An ERROR track has no position_data/error_data (enforce_measurement_backed_verdict's
    own precondition / _create_failed_track never sets them), so the function's
    existing empty-data guard already renders a placeholder instead of reaching the
    sigma_pass-based status_str logic."""
    from laser_trim_analyzer.gui.pages.analyze import AnalyzePage

    class _ChartStub:
        def __init__(self):
            self.placeholder_calls = []

        def show_placeholder(self, msg):
            self.placeholder_calls.append(msg)

    class _PageStub:
        def __init__(self):
            self.chart = _ChartStub()
            self.current_result = None

        def _ensure_chart_initialized(self):
            pass

    stub = _PageStub()
    AnalyzePage._display_track_chart(stub, _error_track())  # must not raise
    assert stub.chart.placeholder_calls  # rendered a placeholder, not a crash


def test_analyze_page_plot_error_vs_position_export_error_track_does_not_crash():
    from laser_trim_analyzer.gui.pages.analyze import AnalyzePage
    from matplotlib.figure import Figure

    fig = Figure()
    ax = fig.add_subplot(111)
    AnalyzePage._plot_error_vs_position_export(None, ax, _error_track())  # must not raise


def test_analyze_page_plot_metrics_summary_error_track_is_not_graded():
    from laser_trim_analyzer.gui.pages.analyze import AnalyzePage
    from matplotlib.figure import Figure

    fig = Figure()
    ax = fig.add_subplot(111)
    AnalyzePage._plot_metrics_summary(None, ax, _error_track())  # must not raise

    joined = "\n".join(t.get_text() for t in ax.texts)
    assert "999.999" not in joined
    assert "not graded" in joined.lower() or "NOT GRADED" in joined
    assert "Sigma Pass" not in joined
    assert "Linearity Pass" not in joined


def test_analyze_page_plot_status_display_error_track_is_not_graded():
    from laser_trim_analyzer.gui.pages.analyze import AnalyzePage
    from matplotlib.figure import Figure

    fig = Figure()
    ax = fig.add_subplot(111)
    AnalyzePage._plot_status_display(None, ax, _error_track(), result=None)  # must not raise

    joined = "\n".join(t.get_text() for t in ax.texts)
    assert "STATUS: FAIL" not in joined
    assert "STATUS: PASS" not in joined
    assert "STATUS: NOT GRADED" in joined.upper()


def test_analyze_page_plot_status_display_graded_tracks_still_show_pass_and_fail():
    """Regression guard for the two branches either side of the new one."""
    from laser_trim_analyzer.gui.pages.analyze import AnalyzePage
    from matplotlib.figure import Figure

    fail_track = _graded_track(AnalysisStatus.FAIL, 0.05, 0.03, False, 0.02, False, fail_points=3)
    fig = Figure()
    ax = fig.add_subplot(111)
    AnalyzePage._plot_status_display(None, ax, fail_track, result=None)
    assert "STATUS: FAIL" in "\n".join(t.get_text() for t in ax.texts)

    pass_track = _graded_track(AnalysisStatus.PASS, 0.01, 0.03, True, 0.005, True)
    fig2 = Figure()
    ax2 = fig2.add_subplot(111)
    AnalyzePage._plot_status_display(None, ax2, pass_track, result=None)
    assert "STATUS: PASS" in "\n".join(t.get_text() for t in ax2.texts)


# ---------------------------------------------------------------------------
# gui/pages/export.py (V5 chart export -- PNG/PDF)
# ---------------------------------------------------------------------------

def test_export_page_plot_error_vs_position_export_error_track_does_not_crash():
    from laser_trim_analyzer.gui.pages.export import ExportPage
    from matplotlib.figure import Figure

    fig = Figure()
    ax = fig.add_subplot(111)
    ExportPage._plot_error_vs_position_export(None, ax, _error_track())  # must not raise


def test_export_page_plot_metrics_summary_error_track_is_not_graded():
    from laser_trim_analyzer.gui.pages.export import ExportPage
    from matplotlib.figure import Figure

    fig = Figure()
    ax = fig.add_subplot(111)
    ExportPage._plot_metrics_summary(None, ax, _error_track())  # must not raise

    joined = "\n".join(t.get_text() for t in ax.texts)
    assert "999.999" not in joined
    assert "not graded" in joined.lower() or "NOT GRADED" in joined
    assert "Sigma Pass" not in joined
    assert "Linearity Pass" not in joined


def test_export_page_plot_status_display_error_track_is_not_graded():
    from laser_trim_analyzer.gui.pages.export import ExportPage
    from matplotlib.figure import Figure

    fig = Figure()
    ax = fig.add_subplot(111)
    ExportPage._plot_status_display(None, ax, _error_track(), result=None)  # must not raise

    joined = "\n".join(t.get_text() for t in ax.texts)
    assert "STATUS: FAIL" not in joined
    assert "STATUS: PASS" not in joined
    assert "STATUS: NOT GRADED" in joined.upper()


def test_export_page_plot_status_display_graded_tracks_still_show_pass_and_fail():
    from laser_trim_analyzer.gui.pages.export import ExportPage
    from matplotlib.figure import Figure

    fail_track = _graded_track(AnalysisStatus.FAIL, 0.05, 0.03, False, 0.02, False, fail_points=3)
    fig = Figure()
    ax = fig.add_subplot(111)
    ExportPage._plot_status_display(None, ax, fail_track, result=None)
    assert "STATUS: FAIL" in "\n".join(t.get_text() for t in ax.texts)

    pass_track = _graded_track(AnalysisStatus.PASS, 0.01, 0.03, True, 0.005, True)
    fig2 = Figure()
    ax2 = fig2.add_subplot(111)
    ExportPage._plot_status_display(None, ax2, pass_track, result=None)
    assert "STATUS: PASS" in "\n".join(t.get_text() for t in ax2.texts)


# ---------------------------------------------------------------------------
# export/excel.py (Excel export)
# ---------------------------------------------------------------------------

def test_excel_export_single_result_error_track_is_not_graded(db, tmp_path):
    from laser_trim_analyzer.export.excel import export_single_result
    import openpyxl

    error_id = db.save_analysis(_wrap(_error_track(), AnalysisStatus.ERROR, "ERR-4"))
    analysis = db.get_analysis(error_id)

    out = tmp_path / "error_track.xlsx"
    export_single_result(analysis, out)  # must not raise

    ws = openpyxl.load_workbook(out)["Track Details"]
    values = [c.value for row in ws.iter_rows() for c in row if c.value is not None]

    assert not any(v == 999.999 for v in values if isinstance(v, (int, float)))
    assert not any(isinstance(v, str) and v.strip().upper() == "NO" for v in values)
    assert any(isinstance(v, str) and "not graded" in v.lower() for v in values)


def test_excel_export_batch_results_error_track_renders_as_dash_not_fail(db, tmp_path):
    """The brief's ~1773 reference (analyze.py::_export_model_results) fetches a
    model's history via get_historical_data() and hands it straight to
    export_batch_results() -> _create_batch_summary_sheet / _create_all_results_sheet.
    Both were already None-safe (render "—") before this fix; pin that with a real
    mixed batch rather than relying on reading the code."""
    from laser_trim_analyzer.export.excel import export_batch_results
    import openpyxl

    db.save_analysis(_wrap(_error_track(), AnalysisStatus.ERROR, "ERR-5"))
    db.save_analysis(_wrap(
        _graded_track(AnalysisStatus.FAIL, 0.05, 0.03, False, 0.02, False, fail_points=3),
        AnalysisStatus.FAIL, "FAIL-5"))
    results = db.get_historical_data(model=MODEL, days_back=3650, limit=10)
    assert len(results) == 2

    out = tmp_path / "batch.xlsx"
    export_batch_results(results, out)  # must not raise

    ws = openpyxl.load_workbook(out)["All Results"]
    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
    by_serial = {r[1]: r for r in rows}  # column B = Serial

    err_row = by_serial["ERR-5"]
    assert not any(v == 999.999 for v in err_row if isinstance(v, (int, float)))
    assert "PASS" not in err_row and "FAIL" not in err_row
    assert err_row.count("—") >= 2  # Sigma Pass and Linearity Pass columns

    fail_row = by_serial["FAIL-5"]
    assert "FAIL" in fail_row
    assert any(isinstance(v, (int, float)) and abs(v - 0.05) < 1e-9 for v in fail_row)


def test_excel_export_single_result_graded_fail_track_still_shows_numbers(db, tmp_path):
    from laser_trim_analyzer.export.excel import export_single_result
    import openpyxl

    fail_id = db.save_analysis(_wrap(
        _graded_track(AnalysisStatus.FAIL, 0.05, 0.03, False, 0.02, False, fail_points=3),
        AnalysisStatus.FAIL, "FAIL-4"))
    analysis = db.get_analysis(fail_id)

    out = tmp_path / "fail_track.xlsx"
    export_single_result(analysis, out)

    ws = openpyxl.load_workbook(out)["Track Details"]
    values = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
    assert any(isinstance(v, str) and "0.050000" in v for v in values)
    assert any(isinstance(v, str) and v.strip().upper() == "NO" for v in values)  # Sigma/Linearity Pass: NO
