"""Model-spec resolution: ONE rule, for the database and for the snapshot a worker carries.

The ingest's analysis looks a model spec up per file -- `get_model_spec`, or `resolve_spec_for_ft`
for a final test, whose serial may name a section. A worker PROCESS never opens a database (ingest-
speed ruling 15), so it answers the same questions from a `SpecSnapshot` taken at folder start
(ruling 16). The two must never answer differently, so they share the rule itself -- the functions
below -- and differ only in where the rows come from: `DatabaseManager` hands them the table's rows
through two queries, the snapshot from memory.

One detail makes that sharing load-bearing rather than tidy: an alias is searched in the TABLE'S
ROW ORDER (id order), so when two specs list the same alias the older row answers. A copy of the
rule that searched in model order answered differently whenever that happened.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from typing import Any, Callable, Dict, Iterable, List, Optional, Tuple

# SpecsMixin (below) needs these three names. `logger` comes from manager.py so its records
# carry the same logger name as before the move (C2 Task 5, step 3) -- same reasoning as
# migrations.py's `from ...manager import logger`. `ModelSpec` is the ORM model the mixin's
# methods query directly. `_specs` is THIS module, imported under the name its moved methods
# already call it by (`_specs.resolve_model_spec(...)` etc., unchanged by the move) -- safe to
# self-import: by the time any SpecsMixin method actually runs, this module has long finished
# loading. The cross-module half (`from ...manager import logger`) is safe for the same reason
# migrations.py's is: manager.py's own (pre-existing) import of THIS module for `_specs` used to
# sit near the top of manager.py, before manager.py's `logger` was defined, which would have
# made this circular -- so that pre-existing import was removed there (it had no other reader
# left once these methods moved) and replaced with one import of `SpecsMixin`, positioned after
# manager.py's `logger`/`ModelSpec`-worthy names are already bound. See the comment at that
# import site in manager.py.
from laser_trim_analyzer.database.manager import logger
from laser_trim_analyzer.database.models import ModelSpec
from laser_trim_analyzer.database import specs as _specs

# A trailing section letter on a final test's serial ('31B', '1004a', '31B '): 8508's sections are
# stored as 8508-A .. 8508-D, while its final-test files say model 8508 and carry the section in
# the serial.
_SECTION_LETTER = re.compile(r'^.*?([A-Za-z])\s*$')


def parse_aliases(aliases_str: Optional[str]) -> List[str]:
    """Pipe-separated aliases as a trimmed list of non-empty tokens."""
    if not aliases_str:
        return []
    return [a.strip() for a in aliases_str.split("|") if a.strip()]


def resolve_model_spec(model: Optional[str],
                       primary: Callable[[str], Optional[Dict[str, Any]]],
                       alias_rows: Callable[[], Iterable[Dict[str, Any]]]
                       ) -> Optional[Dict[str, Any]]:
    """THE model-spec rule (`get_model_spec` and `SpecSnapshot.get_model_spec` alike).

    The model, stripped, matched exactly against the spec rows' `model` (`primary`); failing that,
    the first row -- in id order (`alias_rows`) -- whose pipe-separated aliases contain it
    exactly. None if neither, or for no model at all.
    """
    if not model:
        return None
    model = model.strip()
    row = primary(model)
    if row is not None:
        return row
    for row in alias_rows():
        if model in parse_aliases(row.get("aliases")):
            return row
    return None


def resolve_ft_spec(model: Optional[str], serial: Optional[str],
                    lookup: Callable[[str], Optional[Dict[str, Any]]]
                    ) -> Optional[Dict[str, Any]]:
    """THE final-test rule (`resolve_spec_for_ft` and `SpecSnapshot.resolve_spec_for_ft` alike).

    A serial ending in a letter asks for the section's spec first (`8508` + `31b` -> `8508-B`,
    upper-cased); failing that, or without one, the model's own. `lookup` is the matching
    `get_model_spec`.
    """
    if not model:
        return None
    if serial:
        m = _SECTION_LETTER.match(str(serial))
        if m:
            section_spec = lookup(f"{model}-{m.group(1).upper()}")
            if section_spec:
                return section_spec
    return lookup(model)


@dataclass(frozen=True)
class SpecSnapshot:
    """What the ingest's analysis reads from the database, read ONCE at folder start.

    `specs`: every model_specs row, as `get_model_spec` returns it, kept in id order.
    `ml_thresholds` (model -> sigma threshold) and `ml_predictors` (model -> trained
    ModelPredictor): the ML state a Processor otherwise loads from the database when it is built.

    Plain data -- dicts of plain values, and predictors that pickle -- so it can go to a spawned
    worker process (ingest-speed Task 11). It answers `get_model_spec` and `resolve_spec_for_ft`
    through the same functions `DatabaseManager` does, so it and the database cannot disagree; a
    spec edited while a folder runs reaches the next folder's snapshot.
    """
    specs: Tuple[Dict[str, Any], ...] = ()
    ml_thresholds: Dict[str, float] = field(default_factory=dict)
    ml_predictors: Dict[str, Any] = field(default_factory=dict)
    _by_model: Dict[str, Dict[str, Any]] = field(init=False, repr=False, compare=False)
    _alias_rows: Tuple[Dict[str, Any], ...] = field(init=False, repr=False, compare=False)

    def __post_init__(self):
        rows = tuple(sorted((dict(r) for r in self.specs), key=lambda r: r["id"]))
        object.__setattr__(self, "specs", rows)
        object.__setattr__(self, "_by_model", {r["model"]: r for r in rows})
        # The database's alias query: `aliases IS NOT NULL AND aliases != ''`, in id order.
        object.__setattr__(self, "_alias_rows", tuple(r for r in rows if r.get("aliases")))

    def get_model_spec(self, model: Optional[str]) -> Optional[Dict[str, Any]]:
        row = resolve_model_spec(model, self._by_model.get, lambda: self._alias_rows)
        return None if row is None else dict(row)      # a copy, like the database's own

    def resolve_spec_for_ft(self, model: Optional[str],
                            serial: Optional[str]) -> Optional[Dict[str, Any]]:
        return resolve_ft_spec(model, serial, self.get_model_spec)


class SpecsMixin:
    """Model-spec CRUD: readers, writers, the Excel importer and its angle-text parsers.

    Moved out of `database/manager.py` (2026-09-25, C2 Task 5, step 3) -- a pure move, byte-
    identical method bodies (see the AST proof in this step's commit). Inherited by
    `DatabaseManager`, alongside `MigrationsMixin`, so every existing call site
    (`db.get_model_spec(...)`, `db.save_model_spec(...)`, etc.) is unchanged.

    `get_model_spec` / `resolve_spec_for_ft` still answer through this module's own
    `resolve_model_spec` / `resolve_ft_spec` (via the `_specs` self-import above) -- the ONE rule
    a worker's `SpecSnapshot` (also in this file) answers through as well, so the two can never
    disagree (ingest-speed ruling 16). That sharing did not change: these methods always called
    the same module-level functions this file already defined: joining this file only makes the
    "one rule, one place" arrangement the module docstring describes a little more literal.
    """

    @staticmethod
    def _spec_to_dict(s: "ModelSpec") -> Dict[str, Any]:
        return {
            "id": s.id,
            "model": s.model,
            "element_type": s.element_type,
            "product_class": s.product_class,
            "linearity_type": s.linearity_type,
            "linearity_spec_text": s.linearity_spec_text,
            "linearity_spec_pct": s.linearity_spec_pct,
            "total_resistance_min": s.total_resistance_min,
            "total_resistance_max": s.total_resistance_max,
            "electrical_angle": s.electrical_angle,
            "electrical_angle_tol": s.electrical_angle_tol,
            "electrical_angle_tol_type": getattr(s, "electrical_angle_tol_type", None),
            "electrical_angle_unit": s.electrical_angle_unit,
            "output_smoothness": s.output_smoothness,
            "circuit_type": s.circuit_type,
            "open_closed": getattr(s, "open_closed", None) or s.circuit_type,
            "aliases": getattr(s, "aliases", None),
            "exclude_points": getattr(s, "exclude_points", None),
            "exclude_points_ft": getattr(s, "exclude_points_ft", None),
            "notes": s.notes,
        }

    @staticmethod
    def _parse_aliases(aliases_str: Optional[str]) -> List[str]:
        """Parse pipe-separated aliases into a trimmed list of non-empty tokens."""
        return _specs.parse_aliases(aliases_str)

    def get_all_model_specs(self) -> List[Dict[str, Any]]:
        """Get all model specs as dicts."""
        with self.session() as session:
            specs = session.query(ModelSpec).order_by(ModelSpec.model).all()
            return [self._spec_to_dict(s) for s in specs]

    def get_model_spec(self, model: str) -> Optional[Dict[str, Any]]:
        """
        Get spec for a specific model. Checks both the primary `model` column
        and the pipe-separated `aliases` column, so `1621501` and `2001621501`
        can share a single spec row.

        Answered through `database.specs.resolve_model_spec` -- the ONE rule the
        SpecSnapshot a worker carries answers through as well, so the two can
        never disagree (ingest-speed ruling 16). The rows come from here: the
        exact `model` match, else the rows carrying aliases IN ID ORDER (the order
        the old unordered query read them in, now said out loud -- it decides
        which spec answers an alias two specs share).
        """
        if not model:
            return None
        with self.session() as session:
            def primary(m: str) -> Optional[Dict[str, Any]]:
                spec = session.query(ModelSpec).filter(ModelSpec.model == m).first()
                return None if spec is None else self._spec_to_dict(spec)

            def alias_rows() -> List[Dict[str, Any]]:
                return [self._spec_to_dict(c) for c in session.query(ModelSpec).filter(
                    ModelSpec.aliases.isnot(None),
                    ModelSpec.aliases != "",
                ).order_by(ModelSpec.id)]

            return _specs.resolve_model_spec(model, primary, alias_rows)

    def resolve_spec_for_ft(self, model: Optional[str], serial: Optional[str]) -> Optional[Dict[str, Any]]:
        """
        Resolve a model spec for a Final Test record.

        Multi-section parts (e.g. 8508) store their spec as per-section rows:
        8508-A, 8508-B, 8508-C, 8508-D. But FT files for the same product are
        labeled as model='8508' with the section baked into the serial by the
        operator (e.g. serial='31B' for section B, SN31). This helper tries
        the section-specific spec first, then falls back to the plain model.

        Resolution order:
          1. If serial ends in a letter AND get_model_spec(model-letter) exists,
             return that row.
          2. Otherwise return get_model_spec(model).

        The rule is `database.specs.resolve_ft_spec`, shared with the SpecSnapshot.
        """
        return _specs.resolve_ft_spec(model, serial, self.get_model_spec)

    def save_model_spec(self, data: Dict[str, Any]) -> Tuple[int, bool]:
        """Create or update a model spec. Returns (spec_id, was_update)."""
        with self._write_lock:
            with self.session() as session:
                existing = session.query(ModelSpec).filter(
                    ModelSpec.model == data["model"]
                ).first()

                if existing:
                    for key, value in data.items():
                        if key not in ("id", "model", "created_at", "updated_at"):
                            setattr(existing, key, value)
                    # updated_at handled automatically by onupdate=utc_now
                    session.flush()
                    return existing.id, True
                else:
                    spec = ModelSpec(**{k: v for k, v in data.items() if k != "id"})
                    session.add(spec)
                    session.flush()
                    return spec.id, False

    def delete_model_spec(self, model: str) -> bool:
        """Delete a model spec. Returns True if found and deleted."""
        with self._write_lock:
            with self.session() as session:
                spec = session.query(ModelSpec).filter(
                    ModelSpec.model == model
                ).first()
                if spec:
                    session.delete(spec)
                    return True
                return False

    def get_distinct_element_types(self) -> List[str]:
        """Get all distinct element types from model_specs."""
        with self.session() as session:
            results = session.query(ModelSpec.element_type).filter(
                ModelSpec.element_type.isnot(None)
            ).distinct().order_by(ModelSpec.element_type).all()
            return [r[0] for r in results]

    def get_distinct_product_classes(self) -> List[str]:
        """Get all distinct product classes from model_specs."""
        with self.session() as session:
            results = session.query(ModelSpec.product_class).filter(
                ModelSpec.product_class.isnot(None)
            ).distinct().order_by(ModelSpec.product_class).all()
            return [r[0] for r in results]

    @staticmethod
    def _parse_angle_string(angle_text: Optional[str]) -> Tuple[Optional[float], Optional[float], Optional[str], Optional[str]]:
        """
        Parse a single angle-spec string into (value, tol, unit, tol_type).

        Handles many formats:
          '1.31" ± .005"'        symmetric tolerance
          '.665" +/-.005"'       symmetric tolerance
          '150° ± 1°'            symmetric tolerance
          '350° Min'             one-sided (floor; slope may go up)
          '340° Max'             one-sided (ceiling; slope may go down)
          '89° - 91°'            range (midpoint ± half-range)
          '2.812" - 2.832"'      range
          '±45°', '+/- 27.5°'    bilateral (±N from center)
          '120°', '1.25"'        nominal only, no tolerance
          'See ATP-10312-DS'     reference doc — returns all Nones
          'SEE CHARTS'           reference doc — returns all Nones

        Returns (angle_val, angle_tol, angle_unit, angle_tol_type).
        All None if the text is empty or a reference-doc string.
        """
        import re as _re

        angle_val = None
        angle_tol = None
        angle_unit = None
        angle_tol_type = None

        if not angle_text:
            return angle_val, angle_tol, angle_unit, angle_tol_type

        txt = angle_text.strip()
        if not txt:
            return angle_val, angle_tol, angle_unit, angle_tol_type

        txt_lower = txt.lower()

        # Reference-doc strings: store nothing (don't pull a part
        # number out of the string and call it an angle).
        if (txt_lower.startswith("see ") or
            "see chart" in txt_lower or
            "see table" in txt_lower or
            "see atp" in txt_lower):
            return None, None, None, None

        has_deg = '°' in txt or 'deg' in txt_lower
        has_inch = '"' in txt
        unit_guess = "deg" if has_deg else ("in" if has_inch else None)

        # Bilateral: starts with ± or +/- (e.g. '±45°', '+/- 27.5°')
        bi_match = _re.match(r'^\s*(?:[±]|\+/?-)\s*([\d.]+)', txt)

        # "Min" or "Max" qualifier anywhere in the text.
        has_min = bool(_re.search(r'\bmin\b', txt_lower))
        has_max = bool(_re.search(r'\bmax\b', txt_lower))

        # Range form: "89° - 91°" or "2.812" - 2.832""
        range_match = _re.search(r'([\d.]+)[°"]?\s*[-–]\s*([\d.]+)', txt)

        # Symmetric form: "N ± M" or "N +/- M"
        sym_match = _re.search(r'([\d.]+)[°"]?\s*(?:[±]|\+/?-)\s*([\d.]+)', txt)

        # Priority: symmetric > range > bilateral > min/max > plain
        if sym_match and not (bi_match and bi_match.start() == 0 and '±' not in txt[:3]):
            try:
                angle_val = float(sym_match.group(1))
                angle_tol = float(sym_match.group(2))
                angle_tol_type = "symmetric"
                angle_unit = unit_guess or "in"
            except ValueError:
                pass

        if angle_val is None and range_match:
            try:
                lo = float(range_match.group(1))
                hi = float(range_match.group(2))
                if hi > lo:
                    angle_val = (lo + hi) / 2.0
                    angle_tol = (hi - lo) / 2.0
                    angle_tol_type = "range"
                    angle_unit = unit_guess or "in"
            except ValueError:
                pass

        if angle_val is None and bi_match:
            try:
                angle_val = float(bi_match.group(1))
                angle_tol = None
                angle_tol_type = "bilateral"
                angle_unit = unit_guess or "deg"
            except ValueError:
                pass

        if angle_val is None and has_min:
            num_match = _re.search(r'([\d.]+)', txt)
            if num_match:
                try:
                    angle_val = float(num_match.group(1))
                    angle_tol = None
                    angle_tol_type = "min"
                    angle_unit = unit_guess or "in"
                except ValueError:
                    pass

        if angle_val is None and has_max:
            num_match = _re.search(r'([\d.]+)', txt)
            if num_match:
                try:
                    angle_val = float(num_match.group(1))
                    angle_tol = None
                    angle_tol_type = "max"
                    angle_unit = unit_guess or "in"
                except ValueError:
                    pass

        if angle_val is None:
            num_match = _re.search(r'([\d.]+)', txt)
            if num_match:
                try:
                    angle_val = float(num_match.group(1))
                    angle_tol = None
                    angle_tol_type = None
                    angle_unit = unit_guess or "in"
                except ValueError:
                    pass

        return angle_val, angle_tol, angle_unit, angle_tol_type

    @staticmethod
    def _split_multi_section_angle(angle_text: Optional[str]) -> List[Tuple[List[str], str]]:
        """
        If the angle text describes multiple sections with different specs,
        split it into [(sections, per_section_angle_text), ...].

        Example inputs that trigger splitting:
          'Section A, B & C = 60° +/-.3°\\nSection D = 66.66° +/-.3°'
            -> [(['A','B','C'], '60° +/-.3°'), (['D'], '66.66° +/-.3°')]
          'Sections A, B = 60° ± .3°; Section C = 66° ± .3°'
            -> [(['A','B'], '60° ± .3°'), (['C'], '66° ± .3°')]

        Returns empty list when the text is NOT a multi-section spec — caller
        should then treat the whole string as a single spec.
        """
        import re as _re

        if not angle_text:
            return []

        txt = angle_text.strip()

        # Must contain at least two occurrences of "Section" (case-insensitive)
        # to qualify as multi-section. One "Section X = Y" row is technically
        # possible but pointless to split.
        if len(_re.findall(r'\bsections?\b', txt, _re.IGNORECASE)) < 2:
            return []

        # Split on newlines OR on semicolons — the real-world Excel has
        # '\n' but users may type ';' too.
        raw_parts = [p.strip() for p in _re.split(r'[\n\r;]+', txt) if p.strip()]

        out: List[Tuple[List[str], str]] = []
        for part in raw_parts:
            # Match: "Section(s) A, B & C = <spec text>"
            m = _re.match(
                r'^\s*Sections?\s+([A-Za-z0-9 ,&/]+?)\s*=\s*(.+)$',
                part,
                _re.IGNORECASE,
            )
            if not m:
                continue
            sections_str = m.group(1).strip()
            spec_text = m.group(2).strip()
            # Break 'A, B & C' into ['A','B','C']. Accept ',', '&', ' and '.
            tokens = _re.split(r'[,&/]|\band\b', sections_str, flags=_re.IGNORECASE)
            sections = [t.strip().upper() for t in tokens if t.strip()]
            # Only keep single-letter section labels (A-Z). Drop anything weird
            # to stay conservative.
            sections = [s for s in sections if _re.match(r'^[A-Z]$', s)]
            if sections and spec_text:
                out.append((sections, spec_text))

        return out

    def import_model_specs_from_excel(self, file_path: str) -> Dict[str, int]:
        """
        Import model specs from the reference Excel file.
        Merges: updates existing, adds new, never deletes.

        Returns: {"updated": N, "added": N, "skipped": N}
        """
        import re
        import openpyxl

        wb = openpyxl.load_workbook(file_path, read_only=True)
        result = {"updated": 0, "added": 0, "skipped": 0}

        # Collect data from all three sheets
        model_data = {}  # model -> dict of fields

        # Sheet 1: Model Reference (primary, most complete)
        if "Model Reference" in wb.sheetnames:
            ws = wb["Model Reference"]

            # Detect column positions from header row instead of hardcoding.
            # This handles spreadsheets with or without an extra leading column.
            col_map = {}
            header_aliases = {
                "model": "model",
                "element type": "element_type",
                "linearity": "linearity",
                "total resistance": "resistance",
                "electrical angle": "angle",
                "output smoothness": "smoothness",
                "open/closed": "open_closed",
                "product class": "product_class",
                "aliases": "aliases",
            }
            for header_row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
                if not header_row:
                    break
                for idx, cell in enumerate(header_row):
                    if cell is None:
                        continue
                    key = str(cell).strip().lower()
                    if key in header_aliases:
                        col_map[header_aliases[key]] = idx
            logger.debug(f"Model Reference column map: {col_map}")

            if "model" not in col_map:
                logger.warning("Model Reference sheet has no 'Model' column header — skipping")
            else:
                def _cell(row, field):
                    """Get a cell value by field name, or None if column missing."""
                    idx = col_map.get(field)
                    if idx is None or idx >= len(row) or row[idx] is None:
                        return None
                    return str(row[idx]).strip() or None

                for row in ws.iter_rows(min_row=2, values_only=True):
                    if not row:
                        continue
                    model = _cell(row, "model")
                    if not model:
                        continue

                    element_type = _cell(row, "element_type")
                    linearity_text = _cell(row, "linearity")
                    resistance_text = _cell(row, "resistance")
                    angle_text = _cell(row, "angle")
                    smoothness = _cell(row, "smoothness")
                    open_closed = _cell(row, "open_closed")
                    product_class = _cell(row, "product_class")
                    aliases_raw = _cell(row, "aliases")

                    # Parse linearity type from text
                    linearity_type = None
                    linearity_pct = None
                    if linearity_text:
                        lt_lower = linearity_text.lower()
                        # Extract type: look for (Absolute), (Independent), etc.
                        type_match = re.search(
                            r'\(?(Absolute|Independent|Term Base|Zero-Based|VR Max)\)?',
                            linearity_text, re.IGNORECASE
                        )
                        if type_match:
                            linearity_type = type_match.group(1)
                            # Normalize case
                            type_map = {"absolute": "Absolute", "independent": "Independent",
                                        "term base": "Term Base", "zero-based": "Zero-Based",
                                        "vr max": "VR Max"}
                            linearity_type = type_map.get(linearity_type.lower(), linearity_type)
                        elif any(kw in lt_lower for kw in
                                 ['see chart', 'see table', 'function', 'trim according',
                                  'logarithmic', 'logaithmic', 'bowtie', 'no linearity']):
                            linearity_type = "Custom"

                        # Extract percentage: handle ± N.N%, +/-N.N%, +/-.N%
                        # Try ± first, then +/- variants
                        pct_match = re.search(r'[±]\s*(\d*\.?\d+)\s*%', linearity_text)
                        if not pct_match:
                            pct_match = re.search(r'\+/?-?\s*(\d*\.?\d+)\s*%', linearity_text)
                        if pct_match:
                            try:
                                linearity_pct = float(pct_match.group(1))
                            except ValueError:
                                pass

                    # Parse resistance: "950 - 1,050 Ω" → min=950, max=1050
                    r_min = None
                    r_max = None
                    if resistance_text:
                        r_match = re.search(
                            r'([\d,]+\.?\d*)\s*[-–]\s*([\d,]+\.?\d*)',
                            resistance_text
                        )
                        if r_match:
                            try:
                                r_min = float(r_match.group(1).replace(',', ''))
                                r_max = float(r_match.group(2).replace(',', ''))
                            except ValueError:
                                pass

                    # Parse angle — either a single spec or a multi-section spec.
                    # Multi-section example from Excel (model 8508):
                    #   'Section A, B & C = 60° +/-.3°\nSection D = 66.66° +/-.3°'
                    # In that case we emit one spec row per section letter so the
                    # trim files (which come in as 8508-A, 8508-B, ...) each find
                    # the matching spec via plain model-name lookup.
                    sections = self._split_multi_section_angle(angle_text)

                    # Normalize aliases: accept '|' or ',' as separator, dedupe
                    # and drop empties. Stored as pipe-separated in DB.
                    aliases_norm = None
                    if aliases_raw and aliases_raw not in ("None", "nan"):
                        tokens = re.split(r'[|,]', aliases_raw)
                        clean = []
                        seen = set()
                        for t in tokens:
                            t = t.strip()
                            if t and t not in seen and t != model:
                                seen.add(t)
                                clean.append(t)
                        if clean:
                            aliases_norm = " | ".join(clean)

                    # Shared fields common to every section row for this source row.
                    shared = {
                        "element_type": element_type if element_type and element_type != 'None' else None,
                        "product_class": product_class if product_class and product_class != 'None' else None,
                        "linearity_type": linearity_type,
                        "linearity_spec_text": linearity_text if linearity_text and linearity_text != 'None' else None,
                        "linearity_spec_pct": linearity_pct,
                        "total_resistance_min": r_min,
                        "total_resistance_max": r_max,
                        "output_smoothness": smoothness if smoothness and smoothness != 'None' else None,
                        # Write open_closed to both new and legacy fields so GUIs
                        # reading either column keep working.
                        "open_closed": open_closed if open_closed and open_closed != 'None' else None,
                        "circuit_type": open_closed if open_closed and open_closed != 'None' else None,
                        "aliases": aliases_norm,
                    }

                    if sections:
                        # Multi-section model: emit one row per section letter.
                        for section_letters, per_section_text in sections:
                            angle_val, angle_tol, angle_unit, angle_tol_type = \
                                self._parse_angle_string(per_section_text)
                            for letter in section_letters:
                                section_model = f"{model}-{letter}"
                                model_data[section_model] = {
                                    "model": section_model,
                                    **shared,
                                    "electrical_angle": angle_val,
                                    "electrical_angle_tol": angle_tol,
                                    "electrical_angle_tol_type": angle_tol_type,
                                    "electrical_angle_unit": angle_unit,
                                }
                        logger.info(
                            f"Model specs: expanded {model!r} into "
                            f"{sum(len(s) for s, _ in sections)} section rows"
                        )
                    else:
                        # Normal single-spec row.
                        angle_val, angle_tol, angle_unit, angle_tol_type = \
                            self._parse_angle_string(angle_text)
                        model_data[model] = {
                            "model": model,
                            **shared,
                            "electrical_angle": angle_val,
                            "electrical_angle_tol": angle_tol,
                            "electrical_angle_tol_type": angle_tol_type,
                            "electrical_angle_unit": angle_unit,
                        }

        # Sheet 2: Element Type (supplement — broader coverage)
        if "Element Type" in wb.sheetnames:
            ws = wb["Element Type"]
            for row in ws.iter_rows(min_row=2, values_only=True):
                model = str(row[0]).strip() if row[0] else None
                etype = str(row[1]).strip() if row[1] else None
                if model and etype and etype != 'None':
                    if model not in model_data:
                        model_data[model] = {"model": model, "element_type": etype}
                    elif not model_data[model].get("element_type"):
                        model_data[model]["element_type"] = etype

        # Sheet 3: Product Class (supplement — broadest coverage)
        if "Product Class" in wb.sheetnames:
            ws = wb["Product Class"]
            for row in ws.iter_rows(min_row=2, values_only=True):
                model = str(row[0]).strip() if row[0] else None
                pclass = str(row[1]).strip() if row[1] else None
                if model and pclass and pclass != 'None':
                    if model not in model_data:
                        model_data[model] = {"model": model, "product_class": pclass}
                    elif not model_data[model].get("product_class"):
                        model_data[model]["product_class"] = pclass

        wb.close()

        # Save to database (merge logic — save_model_spec handles upsert atomically)
        for model_name, data in model_data.items():
            try:
                _, was_update = self.save_model_spec(data)
                if was_update:
                    result["updated"] += 1
                else:
                    result["added"] += 1
            except Exception as e:
                logger.warning(f"Skipping model spec {model_name}: {e}")
                result["skipped"] += 1

        logger.info(
            f"Model specs import: {result['added']} added, "
            f"{result['updated']} updated, {result['skipped']} skipped"
        )
        return result
